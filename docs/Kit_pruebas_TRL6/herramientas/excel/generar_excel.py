# Genera el Excel de pruebas con usuarios (UAT y usabilidad, SUS incluido) a partir de un JSON de roles y casos.
# Uso: python3 generar_excel.py <salida.xlsx> [roles.json]      (requiere: pip install openpyxl)
import datetime, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
OUT=sys.argv[1]
F="Arial"
def font(b=False,i=False,c="000000",s=10): return Font(name=F,bold=b,italic=i,color=c,size=s)
YEL=PatternFill("solid",fgColor="FFF4C2"); HEAD=PatternFill("solid",fgColor="1F5F8B")
SUB=PatternFill("solid",fgColor="DCE8F1"); GREY=PatternFill("solid",fgColor="EEEEEE")
thin=Side(style="thin",color="BFBFBF"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
WRAP=Alignment(wrap_text=True,vertical="top"); CEN=Alignment(horizontal="center",vertical="center",wrap_text=True)
import json as _json
_CFG=_json.load(open(sys.argv[2] if len(sys.argv)>2 else "roles.ejemplo.json",encoding="utf-8"))
ROLES=[(r["nombre"],r["prefijo"],[(c["codigo"],c["tarea"],c["esperado"]) for c in r["casos"]],r["participantes_previstos"]) for r in _CFG["roles"]]
NSLOTS=10
SUSQ=["Creo que me gustaría usar esta plataforma con frecuencia.","Encontré la plataforma innecesariamente compleja.","Pensé que la plataforma era fácil de usar.","Creo que necesitaría el apoyo de una persona técnica para poder usar la plataforma.","Encontré que las diferentes funciones de la plataforma estaban bien integradas.","Pensé que había demasiada inconsistencia en la plataforma.","Imagino que la mayoría de las personas aprenderían a usar esta plataforma muy rápido.","Encontré la plataforma muy incómoda de usar.","Me sentí muy seguro/a usando la plataforma.","Necesité aprender muchas cosas antes de poder empezar a usar la plataforma."]
wb=Workbook(); ins=wb.active; ins.title="Instrucciones"; ins.sheet_view.showGridLines=False
ins.column_dimensions["A"].width=3; ins.column_dimensions["B"].width=26; ins.column_dimensions["C"].width=95
r=1
ins.cell(r,2,"Pruebas con usuarios (UAT y usabilidad)").font=font(True,s=14); r+=1
ins.cell(r,2,"Funcionales y de aceptación (UAT) y usabilidad. Una fila por participante en la hoja de su rol.").font=font(i=True,c="555555"); r+=2
def sec(t):
    global r; ins.cell(r,2,t).font=font(True,c="1F5F8B",s=11); r+=1
def row(a,b):
    global r
    ins.cell(r,2,a).font=font(True); ins.cell(r,2).alignment=WRAP
    c=ins.cell(r,3,b); c.font=font(); c.alignment=WRAP; r+=1
sec("Cómo se llena")
row("Celdas amarillas","Son las únicas que se llenan. Lo demás se calcula solo.")
row("Fila gris","Es un ejemplo del formato y no cuenta en el resumen.")
row("Resultado","Lista desplegable: Cumple (sin ayuda) · Con ayuda · No cumple. Solo «Cumple» cuenta como éxito.")
row("Minuto en video","Minuto en que empieza el caso en la grabación, por ejemplo 03:15.")
row("Tiempo (min)","Minutos que tardó la persona en el caso, por ejemplo 2,5.")
row("Errores","Clics equivocados, retrocesos o mensajes de error que vio la persona.")
row("SUS P1 a P10","Respuesta de la persona, del 1 al 5. El puntaje SUS aparece al llenar las 10.")
r+=1
sec("Protocolo de grabación")
row("1. Grabación","Pantalla completa con el audio del micrófono, sin cortes, de principio a fin.")
row("2. Apertura","La persona dice su código, su rol, la fecha y la hora en pantalla, y «Autorizo esta grabación para fines de evaluación».")
row("3. Persona en vivo","El moderador dicta un número al azar; la persona lo dice en voz alta y lo escribe en la plataforma.")
row("4. Durante los casos","Antes de cada caso dice su código (por ejemplo «Caso ADM-01»). Piensa en voz alta. El moderador no ayuda.")
row("5. Casos con efectos","Si un caso genera efectos reales (alertas, avisos, cobros), usar las cuentas del espacio de prueba y avisar antes al equipo.")
row("6. Cierre","Responde las 10 preguntas del SUS en voz alta.")
row("Dónde va el video","evidencia/02-08_usuarios/<rol>/<código>.mp4")
r+=1
sec("Flujo completo (casos marcados [Flujo completo])")
row("Qué demuestra","Una misma entidad de prueba recorre todo el proceso, de punta a punta, con las mismas personas de cada rol.")
row("Orden de las sesiones","<Escriba el orden: por ejemplo 1) Rol 2: R2-01 · 2) Rol 3: R3-01 · 3) Rol 1: R1-01>")
row("Condiciones","Mismo día y misma entidad de prueba (con el prefijo TRL-QA-). Cada sesión empieza cuando la anterior terminó. Anotar en Observaciones el código de la entidad.")
r+=1
sec("Metas")
ins.cell(r,2,"SUS mínimo").font=font(True); ins.cell(r,3,68).font=font(c="0000FF"); ins.cell(r,3).alignment=Alignment(horizontal="left"); MS=f"Instrucciones!$C${r}"; r+=1
ins.cell(r,2,"Éxito mínimo de tareas").font=font(True); c=ins.cell(r,3,0.85); c.font=font(c="0000FF"); c.number_format="0%"; c.alignment=Alignment(horizontal="left"); ME=f"Instrucciones!$C${r}"; r+=1
ins.cell(r,3,"SUS 68 es el promedio estándar de referencia; 85 % de éxito es la meta del plan de usabilidad. Valores en azul, editables.").font=font(i=True,c="555555"); r+=2
sec("Cuestionario SUS (1 = totalmente en desacuerdo · 5 = totalmente de acuerdo)")
for i,q in enumerate(SUSQ,1): row(f"P{i}",q)
row("Cálculo","Impares: respuesta − 1. Pares: 5 − respuesta. Se suman los 10 valores y se multiplica por 2,5. Resultado de 0 a 100.")
summary=[]
for name,pre,cases,minimo in ROLES:
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False
    ws.cell(1,1,f"{name} · casos de prueba").font=font(True,s=13)
    for j,h in enumerate(["Caso","Tarea","Resultado esperado"],1):
        c=ws.cell(3,j,h); c.font=font(True,c="FFFFFF"); c.fill=HEAD
    ws.merge_cells(start_row=3,start_column=3,end_row=3,end_column=12)
    rr=4
    for cid,t,e in cases:
        ws.cell(rr,1,cid).font=font(True); ws.cell(rr,2,t).font=font(); ws.cell(rr,3,e).font=font()
        ws.merge_cells(start_row=rr,start_column=3,end_row=rr,end_column=12)
        ws.cell(rr,2).alignment=WRAP; ws.cell(rr,3).alignment=WRAP; ws.row_dimensions[rr].height=30; rr+=1
    hr=rr+1; sr=hr+1; ex=sr+1; d0=ex+1; d1=d0+NSLOTS-1
    for i,h in enumerate(["Participante (código)","Fecha","Archivo de video"],1): ws.cell(sr,i,h)
    col=4; rescols=[]; errcols=[]
    for cid,_,_ in cases:
        ws.cell(hr,col,cid); ws.merge_cells(start_row=hr,start_column=col,end_row=hr,end_column=col+3)
        for j,h in enumerate(["Resultado","Minuto en video","Tiempo (min)","Errores"]): ws.cell(sr,col+j,h)
        rescols.append(col); errcols.append(col+3); col+=4
    sus0=col; ws.cell(hr,sus0,"SUS (1 a 5)"); ws.merge_cells(start_row=hr,start_column=sus0,end_row=hr,end_column=sus0+9)
    for k in range(10): ws.cell(sr,sus0+k,f"P{k+1}")
    calc0=sus0+10; ws.cell(hr,calc0,"Calculado (no editar)"); ws.merge_cells(start_row=hr,start_column=calc0,end_row=hr,end_column=calc0+5)
    for j,h in enumerate(["Puntaje SUS","Casos intentados","Casos en Cumple","Éxito de tareas","Errores totales","Tiempo promedio por caso (min)"]): ws.cell(sr,calc0+j,h)
    obs=calc0+6; ws.cell(sr,obs,"Observaciones")
    for c in range(1,obs+1):
        for rw in (hr,sr):
            cell=ws.cell(rw,c); cell.border=BOX; cell.alignment=CEN
            cell.fill=HEAD if rw==hr else SUB; cell.font=font(True,c="FFFFFF" if rw==hr else "000000")
    ws.row_dimensions[sr].height=30
    dvr=DataValidation(type="list",formula1='"Cumple,Con ayuda,No cumple"',allow_blank=True); ws.add_data_validation(dvr)
    dvs=DataValidation(type="whole",operator="between",formula1="1",formula2="5",allow_blank=True,error="Respuesta del 1 al 5",showErrorMessage=True); ws.add_data_validation(dvs)
    dvn=DataValidation(type="decimal",operator="greaterThanOrEqual",formula1="0",allow_blank=True); ws.add_data_validation(dvn)
    for rw in range(ex,d1+1):
        isex=rw==ex
        for c in range(1,obs+1):
            cell=ws.cell(rw,c); cell.border=BOX; cell.font=font(i=isex,c="7F7F7F" if isex else "000000")
            if c<calc0 or c==obs: cell.fill=GREY if isex else YEL
        for c in rescols:
            dvr.add(ws.cell(rw,c)); ws.cell(rw,c).alignment=CEN
            dvn.add(ws.cell(rw,c+2)); ws.cell(rw,c+2).number_format="0.0"; dvn.add(ws.cell(rw,c+3))
            ws.cell(rw,c+1).number_format="@"
        for k in range(10): dvs.add(ws.cell(rw,sus0+k)); ws.cell(rw,sus0+k).alignment=CEN
        ws.cell(rw,2).number_format="dd/mm/yyyy"
        P=[f"{L(sus0+k)}{rw}" for k in range(10)]
        ws.cell(rw,calc0,f'=IF(COUNT({P[0]}:{P[9]})=10,(({P[0]}+{P[2]}+{P[4]}+{P[6]}+{P[8]}-5)+(25-{P[1]}-{P[3]}-{P[5]}-{P[7]}-{P[9]}))*2.5,"")').number_format="0.0"
        inten="+".join(f'({L(c)}{rw}<>"")' for c in rescols)
        cum="+".join(f'({L(c)}{rw}="Cumple")' for c in rescols)
        ws.cell(rw,calc0+1,f'=IF($A{rw}="","",{inten})')
        ws.cell(rw,calc0+2,f'=IF($A{rw}="","",{cum})')
        ws.cell(rw,calc0+3,f'=IF($A{rw}="","",IF({L(calc0+1)}{rw}=0,"",{L(calc0+2)}{rw}/{L(calc0+1)}{rw}))').number_format="0%"
        ws.cell(rw,calc0+4,f'=IF($A{rw}="","",'+"+".join(f"N({L(c)}{rw})" for c in errcols)+")")
        tcols=",".join(f"{L(c+2)}{rw}" for c in rescols)
        ws.cell(rw,calc0+5,f'=IF($A{rw}="","",IF(COUNT({tcols})=0,"",AVERAGE({tcols})))').number_format="0.0"
        for j in range(6): ws.cell(rw,calc0+j).alignment=CEN
    ws.cell(ex,1,"EJEMPLO"); ws.cell(ex,2,datetime.date(2026,1,15)); ws.cell(ex,3,f"{pre}-01.mp4")
    resx=["Cumple","Con ayuda","Cumple","No cumple","Cumple"]
    for n,c in enumerate(rescols):
        ws.cell(ex,c,resx[n]); ws.cell(ex,c+1,f"{n*3+1:02d}:{(n*17)%60:02d}"); ws.cell(ex,c+2,round(1.5+n*0.7,1)); ws.cell(ex,c+3,n%3)
    for k,v in enumerate([4,2,4,1,4,2,5,1,4,2]): ws.cell(ex,sus0+k,v)
    ws.cell(ex,obs,"Fila de ejemplo: no cuenta en el resumen")
    ws.column_dimensions["A"].width=16; ws.column_dimensions["B"].width=34; ws.column_dimensions["C"].width=18
    for c in range(4,sus0): ws.column_dimensions[L(c)].width=12
    for c in range(sus0,sus0+10): ws.column_dimensions[L(c)].width=5
    for c in range(calc0,calc0+6): ws.column_dimensions[L(c)].width=11
    ws.column_dimensions[L(obs)].width=40
    ws.freeze_panes=ws.cell(ex,4)
    summary.append((name,minimo,d0,d1,calc0))
rs=wb.create_sheet("Resumen",1); rs.sheet_view.showGridLines=False
rs.cell(1,1,"Resumen por rol").font=font(True,s=14)
rs.cell(2,1,"Se calcula solo a partir de las hojas de cada rol. No incluye las filas de ejemplo.").font=font(i=True,c="555555")
hdr=["Rol","Participantes","Participantes previstos","Casos intentados","Casos en Cumple","Éxito de tareas","Errores promedio por persona","SUS promedio","Cumple meta SUS","Cumple meta de éxito","Tiempo promedio por caso (min)"]
for i,h in enumerate(hdr,1):
    c=rs.cell(4,i,h); c.font=font(True,c="FFFFFF"); c.fill=HEAD; c.alignment=CEN; c.border=BOX
rs.row_dimensions[4].height=32
rr=5
for name,minimo,d0,d1,calc0 in summary:
    q=f"'{name}'"; A=f"{q}!$A${d0}:$A${d1}"
    col=lambda k: f"{q}!${L(calc0+k)}${d0}:${L(calc0+k)}${d1}"
    vals=[name,f"=COUNTA({A})",minimo,f"=SUM({col(1)})",f"=SUM({col(2)})",
      f'=IF(D{rr}=0,"",E{rr}/D{rr})',f'=IF(B{rr}=0,"",SUM({col(4)})/B{rr})',
      f'=IF(COUNT({col(0)})=0,"",AVERAGE({col(0)}))',
      f'=IF(H{rr}="","",IF(H{rr}>={MS},"Sí","No"))',f'=IF(F{rr}="","",IF(F{rr}>={ME},"Sí","No"))',
      f'=IF(COUNT({col(5)})=0,"",AVERAGE({col(5)}))']
    for i,v in enumerate(vals,1):
        c=rs.cell(rr,i,v); c.font=font(c="0000FF" if i==3 else "000000"); c.border=BOX
        c.alignment=CEN if i>1 else Alignment(vertical="center")
    rs.cell(rr,6).number_format="0%"; rs.cell(rr,7).number_format="0.0"; rs.cell(rr,8).number_format="0.0"; rs.cell(rr,11).number_format="0.0"; rr+=1
tr=rr; rs.cell(tr,1,"Total")
for i,f in [(2,f"=SUM(B5:B{tr-1})"),(3,f"=SUM(C5:C{tr-1})"),(4,f"=SUM(D5:D{tr-1})"),(5,f"=SUM(E5:E{tr-1})"),(6,f'=IF(D{tr}=0,"",E{tr}/D{tr})')]: rs.cell(tr,i,f)
rs.cell(tr,6).number_format="0%"
for i in range(1,12):
    c=rs.cell(tr,i); c.border=BOX; c.font=font(True); c.fill=SUB
    if i>1: c.alignment=CEN
rs.cell(tr+2,1,"Participantes previstos según el marco TRL 6: al menos 1 por rol, que no haya participado en el desarrollo ni en las pruebas (valores en azul, editables). Renombre las hojas con los roles reales.").font=font(i=True,c="555555")
rs.column_dimensions["A"].width=20
for i in range(2,12): rs.column_dimensions[L(i)].width=14
wb.save(OUT)
