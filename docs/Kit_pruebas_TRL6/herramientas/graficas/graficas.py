"""Gráficas para el reporte y la documentación a partir de los resultados de k6 y del Excel de usuarios.

Requiere: pip install matplotlib openpyxl
Uso:
  python3 graficas.py k6 <series.csv> <salida.png> [titulo]       # p95 y usuarios virtuales por minuto, y % de errores
  python3 graficas.py sus <pruebas-usuarios.xlsx> <salida.png>    # SUS por participante y por rol, con la meta de 68
"""
import csv, gzip, sys, statistics
from collections import defaultdict
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

def abrir(ruta):
    return gzip.open(ruta, "rt") if ruta.endswith(".gz") else open(ruta, newline="")

def k6(ruta, salida, titulo="Carga"):
    dur, err, vus = defaultdict(list), defaultdict(list), defaultdict(list)
    t0 = None
    with abrir(ruta) as f:
        for r in csv.DictReader(f):
            t = int(float(r["timestamp"])); t0 = t if t0 is None else min(t0, t)
            m, v = r["metric_name"], float(r["metric_value"])
            if m == "http_req_duration": dur[t].append(v)
            elif m == "http_req_failed": err[t].append(v)
            elif m == "vus": vus[t].append(v)
    span = max(dur) - t0 if dur else 0
    paso = 60 if span > 600 else 10          # por minuto en pruebas largas, cada 10 s en las cortas
    def por_minuto(d, fn):
        g = defaultdict(list)
        for t, vs in d.items(): g[(t - t0) // paso * paso / 60].extend(vs)
        ks = sorted(g); return ks, [fn(g[k]) for k in ks]
    p95 = lambda a: sorted(a)[min(len(a) - 1, int(len(a) * 0.95))]
    mx, my = por_minuto(dur, p95)
    ex, ey = por_minuto(err, lambda a: 100 * sum(a) / len(a))
    vx, vy = por_minuto(vus, max)
    fig, ax = plt.subplots(2, 1, figsize=(10, 6.5), sharex=True)
    ax[0].plot(mx, my, color="#1F4E79", marker="o" if len(mx) < 30 else None, ms=3, label="p95 (ms)"); ax[0].axhline(2000, color="#B42318", ls="--", lw=1, label="Umbral 2 s")
    ax[0].set_ylabel("p95 (ms)"); ax[0].legend(loc="upper left")
    a2 = ax[0].twinx(); a2.plot(vx, vy, color="#888", lw=1); a2.set_ylabel("Usuarios virtuales", color="#666")
    ax[1].plot(ex, ey, color="#B42318", marker="o" if len(ex) < 30 else None, ms=3); ax[1].axhline(1, color="#B42318", ls="--", lw=1)
    ax[1].set_ylabel("Errores (%)"); ax[1].set_xlabel("Minuto de la prueba")
    fig.suptitle(titulo); fig.tight_layout(); fig.savefig(salida, dpi=150); print("Gráfica:", salida)

def sus(ruta, salida):
    import openpyxl
    wb = openpyxl.load_workbook(ruta, data_only=False)
    datos = []
    for ws in wb.worksheets[2:]:
        hdr = next((r for r in range(1, 40) if any(ws.cell(r, c).value == "P1" for c in range(1, 120))), None)
        if not hdr: continue
        c1 = next(c for c in range(1, 120) if ws.cell(hdr, c).value == "P1")
        for r in range(hdr + 1, hdr + 40):
            cod = ws.cell(r, 1).value
            if not cod or cod == "EJEMPLO": continue
            p = [ws.cell(r, c1 + k).value for k in range(10)]
            if all(isinstance(x, (int, float)) for x in p):
                datos.append((ws.title, cod, ((p[0]+p[2]+p[4]+p[6]+p[8]-5) + (25-p[1]-p[3]-p[5]-p[7]-p[9])) * 2.5))
    if not datos: print("No hay respuestas SUS completas."); return
    fig, ax = plt.subplots(figsize=(10, 4.5))
    roles = list(dict.fromkeys(d[0] for d in datos))
    for i, rol in enumerate(roles):
        ys = [d[2] for d in datos if d[0] == rol]
        ax.scatter([i] * len(ys), ys, color="#1F4E79", alpha=0.7)
        ax.hlines(statistics.mean(ys), i - 0.25, i + 0.25, color="#1F4E79", lw=2)
    ax.axhline(68, color="#B42318", ls="--", lw=1, label="Meta SUS 68")
    ax.set_xticks(range(len(roles))); ax.set_xticklabels(roles); ax.set_ylim(0, 100); ax.set_ylabel("Puntaje SUS")
    ax.set_title(f"SUS por rol · promedio general {statistics.mean(d[2] for d in datos):.1f} (n = {len(datos)})"); ax.legend()
    fig.tight_layout(); fig.savefig(salida, dpi=150); print("Gráfica:", salida)

if __name__ == "__main__":
    if len(sys.argv) < 4: print(__doc__); sys.exit(1)
    {"k6": lambda: k6(sys.argv[2], sys.argv[3], sys.argv[4] if len(sys.argv) > 4 else "Carga"), "sus": lambda: sus(sys.argv[2], sys.argv[3])}[sys.argv[1]]()
